from docx import Document


def extract_docx_text(file_path: str) -> str:
    text_parts = []
    try:
        doc = Document(file_path)

        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate adjacent cells caused by merged table cells
                    deduped = []
                    for cell in row_cells:
                        if not deduped or deduped[-1] != cell:
                            deduped.append(cell)
                    text_parts.append(" | ".join(deduped))

    except Exception as err:
        print(f"docx extraction error for {file_path}: {err}")

    return "\n\n".join(text_parts)